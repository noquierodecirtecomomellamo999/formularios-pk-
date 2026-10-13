"""Exportación de respuestas y resúmenes a Excel."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .almacen import Almacen
from .analisis import resumir
from .resultados import consolidar


def tabla_respuestas(enc, cons) -> pd.DataFrame:
    """Una fila por cuestionario y una columna por respuesta (las opciones múltiples, en columnas aparte)."""
    filas = []
    for c in cons.cuestionarios:
        fila = {
            "Cuestionario": c.id,
            "Nº impreso": c.hoja or "",
            "Incluido": "No (descartado)" if c.id in cons.descartados else "Sí",
            "Páginas que faltan": ", ".join(str(i + 1) for i in c.faltan),
        }
        resp = cons.respuestas[c.id]
        for e in enc.preguntas:
            r = resp[e.id]
            pref = (e.numero + " " if e.numero else "") + e.texto[:60]
            if e.tipo == "unica":
                fila[pref] = e.opciones[r.valor] if r.valor is not None else ""
                for k, info in r.especifique.items():
                    fila[f"{pref} · {e.opciones[k]} (texto)"] = info.get("texto") or (
                        "(escrito, sin transcribir)" if info.get("escrito") else "")
            elif e.tipo == "multiple":
                for k, o in enumerate(e.opciones):
                    fila[f"{pref} · {o}"] = "X" if r.valor and k in r.valor else ""
            elif e.tipo == "escala":
                etiquetas = list(e.valores) + list(e.extras)
                for i, etq in enumerate(e.filas):
                    v = (r.valor or {}).get(i)
                    fila[f"{pref} · {etq}" if etq else pref] = etiquetas[v] if v is not None else ""
            elif e.tipo == "numero":
                fila[pref] = r.valor if r.valor is not None else ""
            elif e.tipo == "abierta":
                fila[pref] = r.texto or ("(escrito, sin transcribir)" if r.escrito else "")
            if r.estado == "dudoso":
                fila.setdefault("Dudas sin revisar", "")
                fila["Dudas sin revisar"] = (fila["Dudas sin revisar"] + f" {e.numero or e.id}").strip()
        filas.append(fila)
    return pd.DataFrame(filas)


def tabla_resumen(resumenes) -> pd.DataFrame:
    filas = []
    for r in resumenes:
        e = r.elemento
        preg = f"{e.numero} {e.texto}".strip()
        if e.tipo in ("unica", "multiple"):
            for f in r.tabla:
                filas.append({"Pregunta": preg, "Respuesta": f.etiqueta, "Personas": f.n,
                              "Porcentaje": round(f.pct, 1) if f.pct is not None else None})
            filas.append({"Pregunta": preg, "Respuesta": "(sin respuesta)", "Personas": r.n_blanco})
        elif e.tipo == "escala":
            for fe in r.filas:
                filas.append({"Pregunta": preg, "Respuesta": fe.etiqueta, "Personas": fe.n,
                              "Media": round(fe.media, 2) if fe.media is not None else None,
                              "Mediana": fe.mediana,
                              f"% ≥ umbral": round(fe.pct_umbral, 1) if fe.pct_umbral is not None else None,
                              "No puntuables (No usa…)": fe.n_extras})
        elif e.tipo == "numero":
            filas.append({"Pregunta": preg, "Respuesta": "valores", "Personas": len(r.numeros),
                          "Media": round(sum(r.numeros) / len(r.numeros), 2) if r.numeros else None})
        elif e.tipo == "abierta":
            filas.append({"Pregunta": preg, "Respuesta": "han escrito", "Personas": r.n_escritos})
        filas.append({"Pregunta": "", "Respuesta": "", "Resultado": ""})
        filas[-2]["Resultado"] = r.frase
    return pd.DataFrame(filas)


def _formatear(hoja, ancho_max: int = 60) -> None:
    cabecera = PatternFill("solid", fgColor="F0EFEC")
    for celda in hoja[1]:
        celda.font = Font(bold=True)
        celda.fill = cabecera
        celda.alignment = Alignment(wrap_text=True, vertical="top")
    for i, columna in enumerate(hoja.columns, start=1):
        largo = max((len(str(c.value)) for c in columna if c.value is not None), default=8)
        hoja.column_dimensions[get_column_letter(i)].width = min(ancho_max, max(10, largo + 2))
    hoja.freeze_panes = "B2"


def exportar_excel(almacen: Almacen, encuesta_id: str, destino: Path) -> Path:
    enc = almacen.cargar(encuesta_id)
    cons = consolidar(almacen, encuesta_id, enc)
    resumenes = resumir(enc, cons)
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(destino, engine="openpyxl") as w:
        tabla_resumen(resumenes).to_excel(w, sheet_name="Resumen", index=False)
        tabla_respuestas(enc, cons).to_excel(w, sheet_name="Respuestas", index=False)
        errores = pd.DataFrame([{"Lote": r.lote, "Página": r.resultado.origen, "Problema": r.resultado.error}
                                for r in cons.errores])
        if not errores.empty:
            errores.to_excel(w, sheet_name="Páginas no leídas", index=False)
        for nombre in w.sheets:
            _formatear(w.sheets[nombre], 70 if nombre == "Resumen" else 40)
    return destino
