"""Importación de cuestionarios desde texto y desde Word."""

import io

from encuestas_omr.importar import desde_docx, desde_texto
from encuestas_omr.modelo import validar

TEXTO = """
CUESTIONARIO DE SATISFACCIÓN     Fecha:
# 1. Bienvenida
1.1. ¿Cómo se siente acogido/a al llegar al Centro?
☐ Muy bien
☐ Bien
☐ Regular
☐ Mal
1.2. ¿Qué tratamientos recibe? (Puede marcar más de una opción)
- Fisioterapia
- Logopedia
- Otro, ¿cuál? __________
4.3. De los siguientes servicios, valore su nivel de satisfacción (0 = nada / 10 = mucho):
- Transporte: 0 1 2 3 4 5 6 7 8 9 10
- Atención sanitaria: 0 1 2 3 4 5 6 7 8 9 10
7. Refiriéndonos a cómo se encuentra este año, señale:
a) Respecto a su estado físico:
- Ha mejorado
- Igual
- Ha empeorado
8. Valore de 0 a 10 los servicios:
TRANSPORTE
0 1 2 3 4 5 6 7 8 9 10
TRATAMIENTOS QUE RECIBE
FISIOTERAPIA
0 1 2 3 4 5 6 7 8 9 10
6.3. Puntúe de 0 a 10 la calidad global del Centro: 0 1 2 3 4 5 6 7 8 9 10
# 7. Sugerencias y mejoras
______________________________________
MUCHAS GRACIAS POR SU TIEMPO
"""


def test_desde_texto():
    enc = desde_texto(TEXTO)
    tipos = [(e.tipo, e.numero) for e in enc.elementos]
    assert tipos[0] == ("abierta", "") or tipos[0][0] in ("abierta", "texto")  # la línea del título
    por_numero = {e.numero: e for e in enc.elementos if e.numero}
    p11 = por_numero["1.1."]
    assert p11.tipo == "unica" and p11.opciones == ["Muy bien", "Bien", "Regular", "Mal"]
    assert p11.favorables == [0, 1]
    p12 = por_numero["1.2."]
    assert p12.tipo == "multiple" and p12.con_texto == [2]
    p43 = por_numero["4.3."]
    assert p43.tipo == "escala" and p43.filas == ["Transporte", "Atención sanitaria"] and p43.umbral == 8
    assert por_numero["7."].tipo == "texto" and por_numero["a)"].tipo == "unica"
    assert por_numero["8."].filas == ["Transporte", "Fisioterapia"]
    assert por_numero["6.3."].tipo == "escala" and por_numero["6.3."].filas == [""]
    assert enc.elementos[-1].tipo == "abierta"
    assert enc.despedida.startswith("MUCHAS GRACIAS")
    assert validar(enc) == []


def test_desde_docx():
    import docx

    doc = docx.Document()
    doc.add_paragraph("Cuestionario de prueba Fecha:")
    doc.add_paragraph("La siguiente encuesta nos ayudará a mejorar los servicios que ofrecemos. Sus respuestas "
                      "serán confidenciales. Marque una sola opción en cada pregunta.")
    doc.add_paragraph("¿Qué tan satisfecho/a está con los servicios?", style="List Number")
    for o in ("☐ Muy satisfecho/a", "☐ Satisfecho/a", "☐ Insatisfecho/a", "☐ Muy insatisfecho/a"):
        doc.add_paragraph(o, style="List Bullet")
    p = doc.add_paragraph("Califique su satisfacción con los siguientes servicios:\n(escala del 1 al 5)",
                          style="List Number")
    assert p is not None
    for s in ("Fisioterapia", "Logopedia"):
        doc.add_paragraph(f"{s}: ☐1 ☐2 ☐3 ☐4 ☐5", style="List Bullet")
    doc.add_paragraph("Comentarios adicionales o sugerencias:", style="List Number")
    buf = io.BytesIO()
    doc.save(buf)
    enc = desde_docx(buf.getvalue())
    assert enc.titulo == "Cuestionario de prueba"
    assert enc.introduccion.startswith("La siguiente encuesta")
    tipos = [e.tipo for e in enc.elementos]
    assert tipos == ["unica", "escala", "abierta"]
    assert enc.elementos[0].opciones[0] == "Muy satisfecho/a"
    assert enc.elementos[1].valores == ["1", "2", "3", "4", "5"]
    assert enc.elementos[1].filas == ["Fisioterapia", "Logopedia"]
    assert [e.numero for e in enc.elementos] == ["1.", "2.", "3."]
