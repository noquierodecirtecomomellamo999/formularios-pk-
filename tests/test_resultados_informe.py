"""Del escaneo al informe: agrupación de páginas, revisiones, análisis, informe PDF y Excel."""

import random

import pytest

from conftest import hojas_rellenas
from encuestas_omr import simulador as sim
from encuestas_omr.analisis import num, pct, resumir
from encuestas_omr.exportar import exportar_excel
from encuestas_omr.informe_pdf import generar_informe
from encuestas_omr.proceso import procesar_lote
from encuestas_omr.resultados import clave_revision, consolidar

EID = "cirenc-familiares-2026"


def _subir(alm, imagenes, tmp_path, nombre="escaneo.pdf"):
    ruta = sim.guardar_pdf(imagenes, tmp_path / nombre, 200)
    lote = alm.crear_lote(EID, [(nombre, ruta.read_bytes())])
    procesar_lote(alm, EID, lote)
    return lote


def test_numeradas_desordenadas_y_duplicadas(almacen, tmp_path):
    m = almacen.maqueta_actual(EID)
    imagenes, verdad = hojas_rellenas(m, almacen.recursos(EID), tmp_path, n=4, primera_hoja=10, semilla=2)
    random.Random(0).shuffle(imagenes)
    _subir(almacen, imagenes + imagenes[:1], tmp_path)  # una página escaneada dos veces
    cons = consolidar(almacen, EID)
    assert sorted(c.hoja for c in cons.cuestionarios) == [10, 11, 12, 13]
    assert all(not c.faltan for c in cons.cuestionarios)
    assert any("dos veces" in a for a in cons.avisos)
    # Las respuestas consolidadas coinciden con lo marcado
    enc = almacen.cargar(EID)
    for c in cons.cuestionarios:
        for e in enc.preguntas:
            r = cons.respuestas[c.id][e.id]
            if e.tipo == "unica":
                marcado = verdad[c.hoja].get(e.id, [])
                assert r.valor == (marcado[0] if marcado else None)


def test_maestra_sin_numerar_por_orden_y_pagina_que_falta(almacen, tmp_path):
    m = almacen.maqueta_actual(EID)
    imagenes, _ = hojas_rellenas(m, almacen.recursos(EID), tmp_path, n=3, primera_hoja=0, semilla=4)
    # Copias de la hoja maestra: el QR lleva el número 0. Se quita la página 2 del segundo cuestionario.
    imagenes = imagenes[:3] + imagenes[4:]
    _subir(almacen, imagenes, tmp_path)
    cons = consolidar(almacen, EID)
    assert len(cons.cuestionarios) == 3
    assert [c.faltan for c in cons.cuestionarios] == [[], [1], []]
    assert any("faltan páginas" in a for a in cons.avisos)


def test_revisiones_y_descartes(almacen, tmp_path):
    m = almacen.maqueta_actual(EID)
    imagenes, _ = hojas_rellenas(m, almacen.recursos(EID), tmp_path, n=2, primera_hoja=1, semilla=8)
    _subir(almacen, imagenes, tmp_path)
    cons = consolidar(almacen, EID)
    c = cons.cuestionarios[0]
    rev = almacen.revisiones(EID)
    rev["campos"][clave_revision(c.id, "p1")] = [4]
    rev["descartados"] = [cons.cuestionarios[1].id]
    almacen.guardar_revisiones(EID, rev)
    cons = consolidar(almacen, EID)
    r = cons.respuestas[c.id]["p1"]
    assert r.valor == 4 and r.estado == "revisado"
    assert len(cons.validos) == 1


def test_analisis_informe_y_excel(almacen, tmp_path):
    m = almacen.maqueta_actual(EID)
    imagenes, _ = hojas_rellenas(m, almacen.recursos(EID), tmp_path, n=6, primera_hoja=1, semilla=9, textos=True)
    _subir(almacen, imagenes, tmp_path)
    enc = almacen.cargar(EID)
    cons = consolidar(almacen, EID, enc)
    resumenes = resumir(enc, cons)
    p1 = next(r for r in resumenes if r.elemento.id == "p1")
    assert p1.n_validas + p1.n_blanco == 6
    assert sum(f.n for f in p1.tabla) == p1.n_validas
    assert p1.frase.startswith("El ")
    escala = next(r for r in resumenes if r.elemento.tipo == "escala")
    assert all(f.media is None or 1 <= f.media <= 5 for f in escala.filas)
    abierta = next(r for r in resumenes if r.elemento.tipo == "abierta")
    assert abierta.n_escritos == 6

    pdf = generar_informe(almacen, EID, tmp_path / "informe.pdf")
    assert pdf.stat().st_size > 50_000
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(str(pdf))
    assert len(doc) >= 5
    doc.close()
    xlsx = exportar_excel(almacen, EID, tmp_path / "resultados.xlsx")
    import openpyxl

    libro = openpyxl.load_workbook(xlsx)
    assert {"Resumen", "Respuestas"} <= set(libro.sheetnames)
    assert libro["Respuestas"].max_row == 7  # cabecera + 6 cuestionarios


@pytest.mark.parametrize("valor,texto", [(8.95, "8,95"), (9.0, "9"), (None, "–")])
def test_formato_numeros(valor, texto):
    assert num(valor) == texto


def test_formato_porcentajes():
    assert pct(86.3636) == "86,4 %"
    assert pct(84.0) == "84 %"
