"""Prueba de humo del programa empaquetado: se ejecuta con el Python incluido en el paquete.

    "dist/Encuestas CIRENC/python/python.exe" instalador/prueba_paquete.py

Recorre todo el proceso real: crear la encuesta, generar el PDF, rellenar y "escanear"
hojas con el simulador, leerlas, consolidar, hacer el informe PDF y el Excel, e importar
un Word. Falla (código distinto de 0) si algo no funciona.
"""

import io
import random
import sys
import tempfile
from pathlib import Path

import numpy as np

# El ._pth del Python embebido añade la carpeta del programa (..) a sys.path
import encuestas_omr  # noqa: E402
from encuestas_omr import simulador as sim  # noqa: E402
from encuestas_omr.almacen import Almacen  # noqa: E402
from encuestas_omr.analisis import resumir  # noqa: E402
from encuestas_omr.exportar import exportar_excel  # noqa: E402
from encuestas_omr.formulario_pdf import generar_pdf  # noqa: E402
from encuestas_omr.importar import desde_docx  # noqa: E402
from encuestas_omr.informe_pdf import generar_informe  # noqa: E402
from encuestas_omr.proceso import procesar_lote  # noqa: E402
from encuestas_omr.resultados import consolidar  # noqa: E402

print("Código del programa en:", Path(encuestas_omr.__file__).parent)
print("Python:", sys.version)

with tempfile.TemporaryDirectory(prefix="prueba_ñ_") as tmp:
    tmp = Path(tmp)
    alm = Almacen(tmp / "Datos de prueba (José)")
    nuevos = alm.instalar_ejemplos()
    assert len(nuevos) == 4, nuevos
    eid = "cirenc-familiares-2026"
    m = alm.maqueta_actual(eid)
    pdf = generar_pdf(m, tmp / "formulario.pdf", alm.recursos(eid), hojas=[1, 2, 3])
    paginas = sim.rasterizar_pdf(pdf, 200)
    rp, rng = random.Random(1), np.random.default_rng(1)
    P = len(m.paginas)
    imagenes, verdad = [], {}
    for k in range(3):
        resp = sim.respuestas_aleatorias(m, rp, prob_blanco=0.0)
        verdad[k + 1] = resp
        imagenes += [sim.escanear(h, rng, angulo=1.0) for h in
                     sim.rellenar(paginas[k * P:(k + 1) * P], m, resp, 200, sim.ESTILOS["boligrafo"], rng)]
    escaneo = sim.guardar_pdf(imagenes, tmp / "escaneo.pdf", 200)
    lote = alm.crear_lote(eid, [(escaneo.name, escaneo.read_bytes())])
    lectura = procesar_lote(alm, eid, lote)
    assert not any(r.error for r in lectura), [r.error for r in lectura]
    mal = 0
    for r in lectura:
        for c in r.campos:
            if c.estado == "ok" and c.seleccion != sorted(verdad[r.hoja].get(c.clave, [])):
                mal += 1
    assert mal == 0, f"{mal} respuestas mal leídas"
    cons = consolidar(alm, eid)
    assert len(cons.validos) == 3
    resumenes = resumir(alm.cargar(eid), cons)
    assert resumenes and resumenes[0].n_total == 3
    informe = generar_informe(alm, eid, tmp / "informe.pdf")
    excel = exportar_excel(alm, eid, tmp / "resultados.xlsx")
    assert informe.stat().st_size > 20_000 and excel.stat().st_size > 5_000

    import docx

    doc = docx.Document()
    doc.add_paragraph("Encuesta de prueba")
    doc.add_paragraph("1. ¿Le gusta?")
    doc.add_paragraph("☐ Sí")
    doc.add_paragraph("☐ No")
    buf = io.BytesIO()
    doc.save(buf)
    assert desde_docx(buf.getvalue()).preguntas[0].opciones == ["Sí", "No"]

print("PRUEBA DEL PAQUETE CORRECTA: lectura, informe, Excel e importación de Word funcionan.")
